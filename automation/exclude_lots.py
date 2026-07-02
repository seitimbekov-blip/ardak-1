"""Автоматизация действия "Исключить" для обычных (не-ЭМ) лотов на Портале
закупок Фонда (zakup.sk.kz).

Основано на em_agent.py (рабочий скрипт для лотов ЭМ) и инструкции по
исключению лота, которая, по подтверждению Ардака, применима и к обычным
лотам - с той разницей, что при поиске лота не нужно фильтровать по
способу закупки "Электронный магазин" (обычные лоты используют другие
способы - ОТ, ЗЦП, ОИ и т.д.).

Режим по умолчанию - DRY-RUN: скрипт доходит до диалога подтверждения
исключения, вводит причину, делает скриншот и НЕ нажимает финальную
красную кнопку "Исключить" - реальных изменений на портале не происходит.
Для реального исключения нужен явный флаг --live и подтверждение "да".

Запуск:
    python automation/exclude_lots.py путь/к/файлу_заливки.xlsx
    python automation/exclude_lots.py путь/к/файлу_заливки.xlsx --live
"""
from __future__ import annotations

import argparse
import asyncio
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook
from playwright.async_api import Page, async_playwright

DEFAULT_EXCLUSION_REASON = "в связи с корректировкой бюджета"
PORTAL_URL = "https://zakup.sk.kz/"


def read_lots_to_exclude(path: str) -> list[dict]:
    """Читает файл заливки и возвращает строки со статусом 'исключить'.

    Заголовок ищется по наличию колонки "Идентификатор из внешней системы",
    т.к. в реальных файлах строк перед заголовком может быть разное
    количество (титул формы, служебные строки и т.д.)."""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    header_row_idx = None
    for i, row in enumerate(rows):
        if row and any(cell and "Идентификатор из внешней системы" in str(cell) for cell in row):
            header_row_idx = i
            break
    if header_row_idx is None:
        raise ValueError("Не найдена строка заголовков (колонка 'Идентификатор из внешней системы')")

    headers = [str(h).strip() if h else "" for h in rows[header_row_idx]]

    lots = []
    for row in rows[header_row_idx + 1:]:
        rec = dict(zip(headers, row))
        ident = str(rec.get("Идентификатор из внешней системы", "") or "").strip()
        action = str(
            rec.get("Тип действия", rec.get("Тип дейстивя", "")) or ""
        ).strip().lower()
        if not ident or action != "исключить":
            continue
        lots.append(rec)

    print(f"Найдено лотов со статусом 'исключить': {len(lots)}")
    return lots


def get_lot_number(lot: dict) -> str:
    """Номер лота в ИСЭЗ/Портале (колонка №) - по нему ищем лот на странице поиска."""
    return str(lot.get("№", "") or "").strip()


def get_reason(lot: dict) -> str:
    reason = str(lot.get("Причина исключения", "") or "").strip()
    return reason or DEFAULT_EXCLUSION_REASON


async def search_lot(page: Page, lot_number: str, method_label: Optional[str] = None) -> None:
    """Ищет лот по номеру пункта плана. В отличие от em_agent.py (лоты ЭМ),
    для обычных лотов способ закупки не фильтруем - оставляем значение по
    умолчанию, если явно не передан method_label."""
    search_input = page.locator("input[placeholder*='Номер пункта плана']")
    if await search_input.count() == 0:
        search_input = page.locator("input").first
    await search_input.fill(lot_number)
    await page.wait_for_timeout(500)

    if method_label:
        sposob_select = page.locator("select").first
        await sposob_select.select_option(label=method_label)
        await page.wait_for_timeout(500)

    find_btn = page.locator("button:has-text('Найти')").first
    await find_btn.click()
    await page.wait_for_timeout(2000)


async def exclude_lot(page: Page, lot: dict, dry_run: bool, screenshot_dir: Path) -> str:
    lot_number = get_lot_number(lot)
    reason = get_reason(lot)
    ident = str(lot.get("Идентификатор из внешней системы", "") or "").strip()

    print(f"  Идентификатор : {ident}")
    print(f"  Номер лота    : {lot_number}")
    print(f"  Причина       : {reason}")

    if not lot_number:
        raise ValueError("Пустой номер лота (колонка №) - не могу найти лот на портале")

    await search_lot(page, lot_number)

    print("  Шаг: нажимаю Действия -> Исключить...")
    actions_btn = page.locator("button:has-text('Действия')").first
    await actions_btn.wait_for(state="visible", timeout=10000)
    await actions_btn.click()
    await page.wait_for_timeout(1000)

    exclude_btn = page.locator("a:has-text('Исключить'), button:has-text('Исключить')").first
    await exclude_btn.wait_for(state="visible", timeout=5000)
    await exclude_btn.click()
    await page.wait_for_timeout(2000)

    print("  Шаг: ввожу причину исключения...")
    reason_input = page.locator(
        "textarea, input[placeholder*='Причина'], input[placeholder*='причина']"
    ).first
    if await reason_input.count() == 0:
        reason_input = page.locator("dialog input, .modal input, .popup input").first
    await reason_input.wait_for(state="visible", timeout=5000)
    await reason_input.fill(reason)
    await page.wait_for_timeout(500)

    confirm_btn = page.locator(
        "button.btn-danger:has-text('Исключить'), button.red:has-text('Исключить')"
    ).first
    if await confirm_btn.count() == 0:
        confirm_btn = page.locator("button:has-text('Исключить')").last
    await confirm_btn.wait_for(state="visible", timeout=5000)

    if dry_run:
        screenshot_dir.mkdir(parents=True, exist_ok=True)
        shot_path = screenshot_dir / f"dry_run_{lot_number.replace('/', '_')}.png"
        await page.screenshot(path=str(shot_path))
        print(f"  СУХОЙ ПРОГОН: лот {lot_number} был бы исключён с причиной '{reason}'.")
        print(f"  Кнопка подтверждения НЕ нажата. Скриншот диалога: {shot_path}")
        await page.keyboard.press("Escape")
        await page.wait_for_timeout(500)
        return "dry_run"

    await confirm_btn.click()
    await page.wait_for_timeout(3000)
    print("  ИСКЛЮЧЕНО успешно")
    return "excluded"


async def run(lots: list[dict], dry_run: bool, screenshot_dir: Path) -> None:
    async with async_playwright() as p:
        browser = await p.chromium.launch(headless=False, slow_mo=500)
        page = await browser.new_page()

        print("Открываю браузер...")
        await page.goto(PORTAL_URL, wait_until="domcontentloaded", timeout=60000)

        print()
        print("====================================================")
        print("  Войдите на портал zakup.sk.kz вручную.")
        print("  Перейдите на страницу Пункты плана нужного года.")
        print("  Затем нажмите Enter здесь.")
        print("====================================================")
        input()

        mode = (
            "СУХОЙ ПРОГОН (dry-run, реальных изменений не будет)"
            if dry_run
            else "БОЕВОЙ РЕЖИМ (реальные изменения на портале!)"
        )
        print(f"Режим: {mode}")
        if not dry_run:
            confirm = input(
                "Вы уверены, что хотите ИСКЛЮЧИТЬ лоты по-настоящему? Введите 'да': "
            )
            if confirm.strip().lower() != "да":
                print("Отменено пользователем.")
                await browser.close()
                return

        results: dict[str, int] = {}
        errors: list[tuple[str, str]] = []

        for i, lot in enumerate(lots, 1):
            ident = str(lot.get("Идентификатор из внешней системы", "") or "").strip()
            lot_number = get_lot_number(lot)
            label = lot_number if lot_number else ident
            print(f"\n[{i}/{len(lots)}] ИСКЛЮЧИТЬ | Лот: {label} | Ид: {ident}")

            try:
                status = await exclude_lot(page, lot, dry_run, screenshot_dir)
                results[status] = results.get(status, 0) + 1
            except Exception as e:  # noqa: BLE001 - логируем и идём дальше, как просил Ардак
                try:
                    await page.screenshot(path=str(screenshot_dir / f"error_{i}.png"))
                except Exception:
                    pass
                error_msg = str(e)
                print(f"  ОШИБКА: {error_msg}")
                print(f"  Скриншот: error_{i}.png")
                errors.append((label, error_msg))
                continue

        print("\n====================================================")
        print(
            f"ГОТОВО. Обработано лотов: {len(lots)}, результат: {results}, "
            f"ошибок: {len(errors)}"
        )
        if errors:
            print("\nЛоты с ошибками (проверьте вручную и исключите либо скорректируйте):")
            for lbl, err in errors:
                print(f"  {lbl} - {err}")
        print("====================================================")
        input("Нажмите Enter для закрытия браузера...")
        await browser.close()


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Исключение лотов (не-ЭМ) на Портале закупок по файлу заливки."
    )
    parser.add_argument("file", help="Файл заливки (.xlsx) с колонкой 'Тип действия'")
    parser.add_argument(
        "--live",
        action="store_true",
        help="Реально нажимать кнопку подтверждения исключения (по умолчанию - сухой прогон).",
    )
    parser.add_argument(
        "--screenshots",
        default="automation/screenshots",
        help="Папка для скриншотов dry-run и ошибок.",
    )
    args = parser.parse_args()

    lots = read_lots_to_exclude(args.file)
    if not lots:
        print("Нет лотов со статусом 'исключить' в файле.")
        return

    asyncio.run(run(lots, dry_run=not args.live, screenshot_dir=Path(args.screenshots)))


if __name__ == "__main__":
    main()
