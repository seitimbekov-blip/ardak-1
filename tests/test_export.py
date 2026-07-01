import io

from openpyxl import load_workbook

from lot_reconciler.compare import reconcile
from lot_reconciler.export import export_em_agent_file, export_full_report
from lot_reconciler.models import LotStatus
from lot_reconciler.portal_loader import load_portal_file
from lot_reconciler.sap_loader import load_sap_file
from tests.test_loaders_and_compare import build_sample_files


def _build_report():
    sap_buffer, portal_buffer = build_sample_files()
    sap_result = load_sap_file(sap_buffer)
    portal_result = load_portal_file(portal_buffer)
    report = reconcile(sap_result, portal_result)
    return report, sap_result


def test_export_full_report_has_expected_sheets_and_highlight():
    report, _ = _build_report()
    buffer = io.BytesIO()
    export_full_report(report, buffer)
    buffer.seek(0)

    wb = load_workbook(buffer)
    assert set(wb.sheetnames) == {"Сводка", "Сверка", "Аномалии"}

    ws = wb["Сверка"]
    header = [c.value for c in ws[1]]
    assert "Номер лота" in header
    assert "Статус" in header

    lot_col = header.index("Номер лота") + 1
    status_col = header.index("Статус") + 1

    row_by_lot = {}
    for row in ws.iter_rows(min_row=2):
        row_by_lot[row[lot_col - 1].value] = row

    added_row = row_by_lot["10/1 Т"]
    assert added_row[status_col - 1].value == LotStatus.ADD.value
    assert added_row[status_col - 1].fill.start_color.rgb in ("00C6EFCE", "C6EFCE")

    removed_row = row_by_lot["10/4 Т"]
    assert removed_row[status_col - 1].value == LotStatus.REMOVE.value

    anomalies_ws = wb["Аномалии"]
    assert anomalies_ws.max_row > 1


def test_export_em_agent_excludes_unchanged_and_sets_action():
    report, sap_result = _build_report()
    buffer = io.BytesIO()
    export_em_agent_file(report, sap_result, buffer)
    buffer.seek(0)

    wb = load_workbook(buffer)
    ws = wb.active

    lot_id_col = sap_result.column_map["lot_id"]
    action_col = sap_result.column_map["action_type_reference"]

    data_rows = list(ws.iter_rows(min_row=len(sap_result.header_rows_values) + 1))
    lots_in_export = {row[lot_id_col - 1].value: row for row in data_rows}

    assert "10/2 Т" not in lots_in_export  # БЕЗ ИЗМЕНЕНИЙ не выгружается

    add_row = lots_in_export["10/1 Т"]
    assert add_row[action_col - 1].value == "добавить"

    change_row = lots_in_export["10/3 Т"]
    assert change_row[action_col - 1].value == "изменить"

    remove_row = lots_in_export["10/4 Т"]
    assert remove_row[action_col - 1].value == "исключить"
