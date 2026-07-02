import io
from typing import Dict, List, Optional, Union

from openpyxl import Workbook

# Совпадает с раскладкой колонок, ожидаемой sap_loader/portal_loader и
# config/field_aliases.json.
SAP_COLUMNS = {
    "lot_id": 1,
    "action_type_reference": 2,
    "month": 3,
    "quantity": 4,
    "unit_price": 5,
    "priority": 6,
    "method": 7,
    "address": 8,
    "sum_with_vat": 9,
    "delivery_terms": 10,
    "payment_terms": 11,
    "sum_no_vat": 31,  # колонка AE
}

SAP_HEADER_TEXT = {
    1: "Идентификатор из внешней системы",
    2: "Тип действия",
    3: "Месяц осуществления закупок",
    4: "Кол-во, объем",
    5: "Маркетинговая цена за единицу",
    6: "Приоритет закупки",
    7: "Способ закупок",
    8: "Адрес поставки товара",
    9: "Сумма, планируемая для закупки ТРУ с НДС, тенге",
    10: "Условия поставки",
    11: "Условия оплаты",
    31: "Сумма, планируемая для закупок ТРУ без НДС, тенге",
}

PORTAL_COLUMNS = {
    "lot_id": 2,
    "month": 3,
    "quantity": 4,
    "unit_price": 5,
    "priority": 6,
    "method": 7,
    "address": 8,
    "sum_with_vat": 9,
    "delivery_terms": 10,
    "payment_terms": 11,
    "sum_no_vat": 19,  # колонка S
}

PORTAL_HEADER_TEXT = {
    1: "№",
    2: "Идентификатор из внешней системы (служебное поле)",
    3: "Срок осуществления закупки",
    4: "Кол-во, объем",
    5: "Маркетинговая цена за единицу",
    6: "Приоритет закупки",
    7: "Способ закупок",
    8: "Адрес поставки товара",
    9: "Сумма, планируемая для закупки ТРУ с НДС, тенге",
    10: "Условия поставки",
    11: "Условия оплаты",
    19: "Сумма, планируемая для закупок ТРУ без НДС, тенге",
}


def build_sap_workbook(rows: List[Dict[str, object]]) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    for col, text in SAP_HEADER_TEXT.items():
        ws.cell(row=6, column=col, value=text)

    row_idx = 8
    for row in rows:
        for field_name, col in SAP_COLUMNS.items():
            ws.cell(row=row_idx, column=col, value=row.get(field_name))
        row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def build_portal_workbook(
    items: List[Union[Dict[str, object], "Separator", "BlankRow"]]
) -> io.BytesIO:
    wb = Workbook()
    ws = wb.active
    for col, text in PORTAL_HEADER_TEXT.items():
        ws.cell(row=9, column=col, value=text)

    row_idx = 12
    for item in items:
        if isinstance(item, Separator):
            ws.cell(row=row_idx, column=1, value=item.text)
        elif isinstance(item, BlankRow):
            pass
        else:
            for field_name, col in PORTAL_COLUMNS.items():
                ws.cell(row=row_idx, column=col, value=item.get(field_name))
        row_idx += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


class Separator:
    def __init__(self, text: str):
        self.text = text


class BlankRow:
    pass
