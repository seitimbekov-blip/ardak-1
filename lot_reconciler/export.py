"""Экспорт результатов сверки: полный отчёт с подсветкой (F8) и файл для
em_agent.py с колонкой "Тип действия" (F9)."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Union

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from lot_reconciler.models import (
    EM_AGENT_ACTION_VALUES,
    LotStatus,
    ReconciliationReport,
)
from lot_reconciler.sap_loader import SapLoadResult

STATUS_FILLS = {
    LotStatus.ADD: PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid"),
    LotStatus.CHANGE: PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid"),
    LotStatus.REMOVE: PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid"),
    LotStatus.UNCHANGED: None,
}

HEADER_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
HEADER_FONT = Font(bold=True)

# Единое обоснование для исключения лота, используется всегда (подтверждено Ардаком).
EXCLUSION_REASON = "в связи с корректировкой бюджета"


def _format_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _format_diffs(diffs) -> str:
    parts = []
    for diff in diffs:
        parts.append(
            f"{diff.label}: было={_format_value(diff.sap_value)} -> "
            f"стало={_format_value(diff.portal_value)}"
        )
    return "; ".join(parts)


def _autosize_columns(ws, max_width: int = 60) -> None:
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            length = len(str(cell.value))
            widths[cell.column] = min(max(widths.get(cell.column, 10), length + 2), max_width)
    for col_idx, width in widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def export_full_report(
    report: ReconciliationReport, output_path: Union[str, Path]
) -> None:
    wb = Workbook()

    summary_ws = wb.active
    summary_ws.title = "Сводка"
    summary_ws.append(["Дата формирования", datetime.now().strftime("%Y-%m-%d %H:%M:%S")])
    summary_ws.append(["Прочитано строк SAP (актуальные+старые версии)", report.sap_row_count])
    summary_ws.append(["Прочитано строк Портала", report.portal_row_count])
    summary_ws.append([])
    summary_ws.append(["Статус", "Количество лотов"])
    for status_name, count in report.counts_by_status().items():
        summary_ws.append([status_name, count])
    summary_ws.append([])
    summary_ws.append(["Аномалий обнаружено", len(report.anomalies)])
    for cell in summary_ws["A5:B5"][0]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    _autosize_columns(summary_ws)

    ws = wb.create_sheet("Сверка")
    headers = [
        "Номер лота",
        "Вид закупки",
        "Статус",
        "Тип действия (SAP, справочно)",
        "Сумма без НДС (SAP)",
        "Сумма без НДС (Портал)",
        "Изменённые поля (было -> стало)",
        "Строка SAP",
        "Строка Портал",
    ]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"

    for result in report.results:
        sap_sum = result.sap_record.fields.get("sum_no_vat") if result.sap_record else None
        portal_sum = (
            result.portal_record.fields.get("sum_no_vat") if result.portal_record else None
        )
        row = [
            result.display_lot_number,
            result.kind or "",
            result.status.value,
            result.action_type_reference or "",
            _format_value(sap_sum),
            _format_value(portal_sum),
            _format_diffs(result.diffs),
            result.sap_record.row_number if result.sap_record else "",
            result.portal_record.row_number if result.portal_record else "",
        ]
        ws.append(row)
        fill = STATUS_FILLS.get(result.status)
        if fill is not None:
            for cell in ws[ws.max_row]:
                cell.fill = fill

    _autosize_columns(ws)

    anomalies_ws = wb.create_sheet("Аномалии")
    anomalies_ws.append(["Тип", "Сообщение", "Лот (базовый номер, вид)", "Строки"])
    for cell in anomalies_ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    for anomaly in report.anomalies:
        lot_display = ""
        if anomaly.lot_key:
            lot_display = f"{anomaly.lot_key[0]} {anomaly.lot_key[1] or ''}".strip()
        anomalies_ws.append(
            [
                anomaly.kind,
                anomaly.message,
                lot_display,
                ", ".join(str(r) for r in anomaly.row_numbers),
            ]
        )
    _autosize_columns(anomalies_ws)

    wb.save(output_path)


def export_em_agent_file(
    report: ReconciliationReport,
    sap_result: SapLoadResult,
    output_path: Union[str, Path],
) -> None:
    """Формирует файл, готовый к загрузке в em_agent.py.

    Для ДОБАВИТЬ/ИЗМЕНИТЬ используется исходная строка SAP (полная структура
    колонок, совпадает с официальным шаблоном загрузочного файла Портала),
    с перезаписанной колонкой "Тип действия" рассчитанным значением. Для
    ИСКЛЮЧИТЬ (лота уже нет в SAP) строка собирается из полей Портала,
    перенесённых в соответствующие колонки SAP по маппингу.

    Колонка "№" (номер лота в ИСЭЗ/Портале) обязательна для ИЗМЕНИТЬ и
    ИСКЛЮЧИТЬ (переносится из Портала) и заполняется "#N/A" для ДОБАВИТЬ
    (лота ещё нет на Портале - подтверждено на реальном образце загрузочного
    файла). Колонка "Причина исключения" для ИСКЛЮЧИТЬ всегда заполняется
    единым обоснованием "в связи с корректировкой бюджета" (подтверждено
    Ардаком - используется всегда, независимо от лота).
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Заливка"

    max_col = sap_result.max_col
    for header_row in sap_result.header_rows_values:
        ws.append(header_row[:max_col])
    for row_cells in ws.iter_rows():
        for cell in row_cells:
            cell.font = HEADER_FONT

    lot_id_col = sap_result.column_map.get("lot_id")
    action_col = sap_result.column_map.get("action_type_reference")
    isez_number_col = sap_result.column_map.get("isez_number")
    exclusion_reason_col = sap_result.column_map.get("exclusion_reason")

    exportable = [r for r in report.results if r.status != LotStatus.UNCHANGED]

    for result in exportable:
        action_value = EM_AGENT_ACTION_VALUES[result.status]

        if result.sap_record is not None:
            row_values = dict(result.sap_record.raw_row)
        else:
            row_values = {col: None for col in range(1, max_col + 1)}
            if lot_id_col:
                row_values[lot_id_col] = result.portal_record.identifier.raw
            if result.portal_record is not None:
                for canonical_field, value in result.portal_record.fields.items():
                    col = sap_result.column_map.get(canonical_field)
                    if col:
                        row_values[col] = value

        if action_col:
            row_values[action_col] = action_value

        if isez_number_col:
            if result.status == LotStatus.ADD:
                row_values[isez_number_col] = "#N/A"
            elif result.portal_record is not None:
                row_values[isez_number_col] = result.portal_record.fields.get("isez_number")

        if result.status == LotStatus.REMOVE and exclusion_reason_col:
            row_values[exclusion_reason_col] = EXCLUSION_REASON

        ws.append([row_values.get(c) for c in range(1, max_col + 1)])

    _autosize_columns(ws)
    wb.save(output_path)
